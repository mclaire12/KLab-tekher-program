"""Extract plain text from uploaded PDF, DOCX and TXT files.

PDF  -> PyMuPDF (``pymupdf``)
DOCX -> python-docx
TXT  -> standard Python decoding (UTF-8 first, then common fallbacks)
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from pathlib import Path

SUPPORTED_TYPES = ("pdf", "docx", "txt")


@dataclass
class ExtractedDocument:
    filename: str
    file_type: str
    size_bytes: int
    text: str
    num_pages: int | None = None  # None when the format has no reliable page count
    warnings: list[str] = field(default_factory=list)

    @property
    def is_empty(self) -> bool:
        return not self.text.strip()


class UnsupportedFileTypeError(ValueError):
    pass


def extract_text(data: bytes, filename: str) -> ExtractedDocument:
    """Extract text from raw file bytes; the format is chosen by file extension."""
    file_type = Path(filename).suffix.lower().lstrip(".")
    if file_type not in SUPPORTED_TYPES:
        raise UnsupportedFileTypeError(
            f"Unsupported file type '.{file_type}'. Supported: {', '.join(SUPPORTED_TYPES)}"
        )

    if file_type == "pdf":
        text, pages, warnings = _extract_pdf(data)
    elif file_type == "docx":
        text, pages, warnings = _extract_docx(data)
    else:
        text, pages, warnings = _extract_txt(data)

    doc = ExtractedDocument(filename, file_type, len(data), text, pages, warnings)
    if doc.is_empty:
        doc.warnings.append(
            "No text could be extracted. The file may be empty or a scanned "
            "image (OCR is not part of this baseline)."
        )
    return doc


def extract_text_from_path(path: str | Path) -> ExtractedDocument:
    path = Path(path)
    return extract_text(path.read_bytes(), path.name)


def _extract_pdf(data: bytes) -> tuple[str, int, list[str]]:
    import pymupdf

    warnings = []
    with pymupdf.open(stream=data, filetype="pdf") as pdf:
        pages = [page.get_text("text") for page in pdf]
    empty_pages = sum(1 for p in pages if not p.strip())
    if empty_pages and empty_pages < len(pages):
        warnings.append(f"{empty_pages} page(s) contained no extractable text.")
    return "\n".join(pages), len(pages), warnings


def _extract_docx(data: bytes) -> tuple[str, None, list[str]]:
    import docx

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    # DOCX files do not store a reliable page count (it depends on rendering).
    return "\n".join(parts), None, []


def _extract_txt(data: bytes) -> tuple[str, None, list[str]]:
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16"), None, []
    try:
        return data.decode("utf-8-sig"), None, []
    except UnicodeDecodeError:
        pass
    # Legacy encodings: tell the user a fallback was used.
    try:
        return data.decode("cp1252"), None, ["File was not UTF-8; decoded as Windows-1252."]
    except UnicodeDecodeError:
        return data.decode("latin-1"), None, ["File was not UTF-8; decoded as Latin-1."]
